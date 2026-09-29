import openpyxl
from pathlib import Path
from django.core.management.base import BaseCommand
from chatbot.models import Business


class Command(BaseCommand):
    help = 'Import business data from DATA-FILEE.xlsx'

    def handle(self, *args, **options):
        xlsx_path = Path(r'C:\Users\rishv\Downloads\DATA-FILEE.xlsx')

        if not xlsx_path.exists():
            self.stdout.write(self.style.ERROR(f'File not found: {xlsx_path}'))
            return

        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
        ws = wb['Business Data']

        headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]

        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get('Business ID'):
                continue

            Business.objects.update_or_create(
                business_id=data['Business ID'],
                defaults={
                    'name': data.get('Business Name', ''),
                    'category': data.get('Category', ''),
                    'subcategory': data.get('Subcategory', ''),
                    'description': data.get('Description', ''),
                    'services': data.get('Services', ''),
                    'address': data.get('Address', ''),
                    'area': data.get('Area', ''),
                    'city': data.get('City', 'Coimbatore'),
                    'state': data.get('State', 'Tamil Nadu'),
                    'pincode': str(data.get('Pincode', '')),
                    'latitude': float(data.get('Latitude', 0)),
                    'longitude': float(data.get('Longitude', 0)),
                    'phone': str(data.get('Phone', '')),
                    'rating': float(data.get('Rating', 0)),
                    'review_count': int(data.get('Review Count', 0)),
                    'price_range': data.get('Price Range', ''),
                    'open_hours': data.get('Open Hours', ''),
                    'search_keywords': data.get('Search Keywords', ''),
                    'status': data.get('Status', 'active'),
                }
            )
            count += 1

        self.stdout.write(self.style.SUCCESS(f'Successfully imported {count} businesses'))
